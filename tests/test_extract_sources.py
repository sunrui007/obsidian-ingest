"""Behavioral extraction checks, using temporary sources and no vault writes."""
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1] / "skills" / "obsidian-ingest"
sys.path.insert(0, str(ROOT / "scripts"))
if (ROOT / ".deps").is_dir():
    sys.path.insert(0, str(ROOT / ".deps"))
import extract_sources as sources


class ExtractTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.work = self.root / "staged"

    def tearDown(self):
        self.temp.cleanup()

    def source(self, name, text):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def assert_assets_resolve(self, result):
        for item in result["assets"]:
            self.assertTrue(Path(item["path"]).is_file())
            self.assertEqual(Path(item["path"]).name, item["name"])
            self.assertTrue(Path(item["path"]).resolve().is_relative_to(self.work.resolve()))

    def test_markdown_preserves_body_title_and_local_image(self):
        image = self.root / "测试 图.png"
        image.write_bytes(b"example-image")
        source = self.source("article.md", "# 正文标题\n\n保留正文。\n\n![图示](<测试 图.png>)")
        result = sources.extract(str(source), self.work)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(result["title"], "正文标题")
        self.assertIn("保留正文。", result["markdown"])
        self.assertIn(result["assets"][0]["name"], result["markdown"])
        self.assert_assets_resolve(result)
        self.assertEqual(source.read_text("utf-8"), "# 正文标题\n\n保留正文。\n\n![图示](<测试 图.png>)")

    def test_legacy_doc_is_blocked_without_binary_guess(self):
        path = self.root / "old.doc"
        path.write_bytes(b"\xd0\xcf\x11\xe0\x00\xffbinary")
        result = sources.extract(str(path), self.work)
        self.assertEqual(result["status"], "blocked")
        self.assertEqual(result["markdown"], "")
        self.assertIn(".docx", result["warnings"][0])

    def test_local_link_cannot_read_outside_source_directory(self):
        outside = self.source("outside.txt", "private-outside-content")
        folder = self.root / "sources"; folder.mkdir()
        source = folder / "article.md"
        source.write_text("# Article\n\n[x](../outside.txt)", encoding="utf-8")
        original_read = Path.read_bytes
        def guarded_read(path):
            self.assertNotEqual(path.resolve(), outside.resolve(), "outside reference must never be read")
            return original_read(path)
        with patch.object(Path, "read_bytes", guarded_read):
            result = sources.extract(str(source), self.work)
        self.assertEqual(result["assets"], [])
        self.assertIn("[x](../outside.txt)", result["markdown"])
        self.assertTrue(any("越出源文档目录" in warning for warning in result["warnings"]))

    def test_linked_markdown_is_not_copied_as_attachment(self):
        self.source("linked.md", "# Another note\nDo not bundle this note.")
        source = self.source("article.md", "# Article\n\n[Another](linked.md)")
        result = sources.extract(str(source), self.work)
        self.assertEqual(result["assets"], [])
        self.assertIn("[Another](linked.md)", result["markdown"])
        self.assertTrue(any("附件类型" in warning for warning in result["warnings"]))

    def test_symlink_attachment_cannot_escape_source_directory(self):
        outside = self.root / "private.png"; outside.write_bytes(b"outside-private-image")
        folder = self.root / "sources"; folder.mkdir()
        link = folder / "linked.png"
        try:
            link.symlink_to(outside)
        except OSError as error:
            self.skipTest(f"This host cannot create symlinks: {error}")
        source = folder / "article.md"; source.write_text("![image](linked.png)", encoding="utf-8")
        original_read = Path.read_bytes
        def guarded_read(path):
            self.assertNotEqual(path.resolve(), outside.resolve(), "symlink target must never be read")
            return original_read(path)
        with patch.object(Path, "read_bytes", guarded_read):
            result = sources.extract(str(source), self.work)
        self.assertEqual(result["assets"], [])
        self.assertEqual(result["markdown"], "![image](linked.png)")
        self.assertTrue(any("越出源文档目录" in warning for warning in result["warnings"]))

    def test_missing_source_is_explicit(self):
        result = sources.extract(str(self.root / "missing.pdf"), self.work)
        self.assertEqual(result["status"], "blocked")
        self.assertIn("不存在", result["warnings"][0])

    @unittest.skipUnless(importlib.util.find_spec("bs4") and importlib.util.find_spec("markdownify"), "HTML dependencies absent")
    def test_wechat_extracts_article_preserving_structure(self):
        html = '''<html><title>wrong title</title><body><h1 id="activity-name">微信文章</h1>
        <nav>导航不要导入</nav><div id="js_content"><h2>第一节</h2><p>'''
        html += "这是微信文章的详细正文。" * 15
        html += '''</p><ul><li>重点一</li></ul><pre><code>print(1)</code></pre>
        <table><tr><th>列名</th></tr><tr><td>单元格</td></tr></table>
        <img data-src="https://mmbiz.qpic.cn/photo.jpg" alt="原图说明"><p>图片注释</p>
        </div><footer>不导入页脚</footer></body></html>'''
        with patch.object(sources, "_fetch", return_value=(html.encode(), "https://mp.weixin.qq.com/s/abc", "text/html")):
            result = sources.extract("https://mp.weixin.qq.com/s/abc", self.work)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(result["title"], "微信文章")
        for expected in ["## 第一节", "- 重点一", "print(1)", "| 列名 |", "https://mmbiz.qpic.cn/photo.jpg", "图片注释"]:
            self.assertIn(expected, result["markdown"])
        self.assertNotIn("导航不要导入", result["markdown"])
        self.assertNotIn("不导入页脚", result["markdown"])

    @unittest.skipUnless(importlib.util.find_spec("bs4") and importlib.util.find_spec("markdownify"), "HTML dependencies absent")
    def test_title_only_and_verification_pages_need_browser(self):
        for html in ["<h1>只有标题</h1>", "<h1>安全验证</h1><p>请完成验证</p>"]:
            with self.subTest(html=html), patch.object(sources, "_fetch", return_value=(html.encode(), "https://example.com/post", "text/html")):
                result = sources.extract("https://example.com/post", self.work)
                self.assertEqual(result["status"], "needs_browser")

    def test_http_failure_is_explicit(self):
        with patch.object(sources, "_fetch", side_effect=RuntimeError("403 Forbidden")):
            result = sources.extract("https://mp.weixin.qq.com/s/abc", self.work)
        self.assertEqual(result["status"], "needs_browser")
        self.assertIn("403", result["warnings"][0])

    @unittest.skipUnless(importlib.util.find_spec("bs4") and importlib.util.find_spec("markdownify"), "HTML dependencies absent")
    def test_actual_http_article_download(self):
        from http.server import BaseHTTPRequestHandler, HTTPServer
        from threading import Thread
        html = ("<html><meta charset='utf-8'><article><h1>本地测试文章</h1><p>" + "HTTP 正文段落。" * 20 + "</p></article></html>").encode()
        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                self.send_response(200); self.send_header("Content-Type", "text/html; charset=utf-8")
                self.end_headers(); self.wfile.write(html)
            def log_message(self, *args): pass
        with HTTPServer(("127.0.0.1", 0), Handler) as server:
            thread = Thread(target=server.handle_request); thread.start()
            result = sources.extract(f"http://127.0.0.1:{server.server_port}/article", self.work)
            thread.join(timeout=5)
        self.assertEqual(result["status"], "complete", result["warnings"])
        self.assertEqual(result["title"], "本地测试文章")
        self.assertIn("HTTP 正文段落。", result["markdown"])

    def test_docx_paragraph_table_and_image_are_preserved(self):
        from docx import Document
        from docx.shared import Inches
        from PIL import Image
        image = self.root / "diagram.png"
        Image.new("RGB", (32, 32), "blue").save(image)
        document = Document()
        document.add_heading("文档标题", 1)
        document.add_paragraph("表格前的段落")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "名称"; table.cell(0, 1).text = "内容"
        table.cell(1, 0).text = "字段"; table.cell(1, 1).text = "保留|竖线"
        document.add_picture(str(image), width=Inches(1))
        document.add_paragraph("表格后的段落")
        path = self.root / "document.docx"; document.save(path)
        result = sources.extract(str(path), self.work)
        self.assertEqual(result["status"], "complete", result["warnings"])
        self.assertIn("# 文档标题", result["markdown"])
        self.assertIn("| 名称 | 内容 |", result["markdown"])
        self.assertIn("保留\\|竖线", result["markdown"])
        self.assertIn("![文档图片]", result["markdown"])
        self.assertLess(result["markdown"].index("表格前"), result["markdown"].index("| 名称"))
        self.assertLess(result["markdown"].index("| 名称"), result["markdown"].index("表格后"))
        self.assertEqual(len(result["assets"]), 2)
        self.assert_assets_resolve(result)

    def test_image_no_ocr_does_not_claim_success(self):
        from PIL import Image
        path = self.root / "image.png"; Image.new("RGB", (64, 64), "white").save(path)
        with patch.object(sources, "_ocr", return_value=("", "OCR unavailable")):
            result = sources.extract(str(path), self.work)
        self.assertEqual(result["status"], "needs_ocr")
        self.assertEqual(result["metadata"]["width"], 64)
        self.assertIn("![原始图片]", result["markdown"])
        self.assert_assets_resolve(result)

    @unittest.skipUnless(importlib.util.find_spec("rapidocr_onnxruntime") and Path("C:/Windows/Fonts/msyh.ttc").is_file(), "OCR dependency or Chinese fixture font absent")
    def test_actual_chinese_image_and_scanned_pdf_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.pdfgen import canvas
        path = self.root / "chinese.png"
        image = Image.new("RGB", (1200, 500), "white")
        draw = ImageDraw.Draw(image); font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 58)
        draw.text((60, 70), "知识库自动分类", font=font, fill="black")
        draw.text((60, 180), "PDF OCR TEST 123", font=font, fill="black")
        image.save(path)
        result = sources.extract(str(path), self.work)
        self.assertEqual(result["status"], "complete", result["warnings"])
        self.assertIn("知识库自动分类", result["markdown"])
        self.assertIn("PDF OCR TEST 123", result["markdown"])
        if importlib.util.find_spec("pymupdf"):
            pdf = self.root / "chinese-scan.pdf"; document = canvas.Canvas(str(pdf))
            document.drawImage(str(path), 0, 180, width=595, height=248); document.save()
            result = sources.extract(str(pdf), self.work)
            self.assertEqual(result["status"], "complete", result["warnings"])
            self.assertEqual(result["metadata"]["pages"], [{"page": 1, "status": "complete", "method": "ocr"}])
            self.assertIn("知识库自动分类", result["markdown"])

    def test_pdf_mixed_text_scan_tracks_each_page(self):
        from reportlab.pdfgen import canvas
        from PIL import Image
        image = self.root / "scan.png"; Image.new("RGB", (700, 900), "gray").save(image)
        path = self.root / "mixed.pdf"
        document = canvas.Canvas(str(path))
        document.drawString(40, 800, "Native text content on page one. " * 3); document.showPage()
        document.drawImage(str(image), 0, 0, width=595, height=842); document.showPage(); document.save()
        with patch.object(sources, "_ocr", return_value=("", "OCR unavailable")):
            result = sources.extract(str(path), self.work)
        self.assertEqual(result["status"], "needs_ocr", result["warnings"])
        self.assertEqual(result["metadata"]["page_count"], 2)
        self.assertEqual([p["status"] for p in result["metadata"]["pages"]], ["complete", "needs_ocr"])
        self.assertIn("Native text content", result["markdown"])
        self.assertGreaterEqual(len(result["assets"]), 2)
        self.assert_assets_resolve(result)

    @unittest.skipUnless(importlib.util.find_spec("pymupdf"), "PyMuPDF absent")
    def test_pdf_scan_with_footer_is_not_mistaken_for_complete_text(self):
        from reportlab.pdfgen import canvas
        from PIL import Image
        image = self.root / "scan.png"; Image.new("RGB", (700, 900), "gray").save(image)
        path = self.root / "footer.pdf"; document = canvas.Canvas(str(path))
        document.drawImage(str(image), 0, 0, width=595, height=842)
        document.drawString(10, 20, "Footer with enough extractable text to hide an otherwise scanned body")
        document.save()
        with patch.object(sources, "_ocr", return_value=("", "OCR unavailable")):
            result = sources.extract(str(path), self.work)
        self.assertEqual(result["status"], "needs_ocr")
        self.assertIn("Footer", result["markdown"])


if __name__ == "__main__":
    unittest.main()
